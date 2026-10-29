"""Revised workbooks preserve prompt and empirical cells outside named updates."""

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]


def test_revised_workbooks_have_only_documented_cell_changes():
    record = json.loads(
        (ROOT / "provenance/workbook_update_20261029.json").read_text(encoding="utf-8")
    )
    for item in record["workbooks"]:
        current_path = ROOT / item["path"]
        original_path = ROOT / item["original_path"]
        assert hashlib.sha256(current_path.read_bytes()).hexdigest() == item["sha256"]
        assert hashlib.sha256(original_path.read_bytes()).hexdigest() == item["original_sha256"]
        current = load_workbook(current_path, read_only=True, data_only=False)
        original = load_workbook(original_path, read_only=True, data_only=False)
        try:
            assert current.sheetnames == original.sheetnames
            observed = set()
            for sheet in current:
                old_sheet = original[sheet.title]
                assert (sheet.max_row, sheet.max_column) == (old_sheet.max_row, old_sheet.max_column)
                for row_index, (new_row, old_row) in enumerate(
                    zip(sheet.values, old_sheet.values), start=1
                ):
                    assert len(new_row) == len(old_row)
                    for col_index, (new, old) in enumerate(zip(new_row, old_row), start=1):
                        if new != old:
                            observed.add((sheet.title, f"{get_column_letter(col_index)}{row_index}"))
            expected = {(cell["sheet"], cell["cell"]) for cell in item["changed_cells"]}
            assert observed == expected
        finally:
            current.close()
            original.close()
